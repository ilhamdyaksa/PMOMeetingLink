import streamlit as st
import pandas as pd
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import io
import datetime
import zipfile
import re
import os
import glob

# ==============================================================================
# KONFIGURASI HALAMAN
# ==============================================================================
st.set_page_config(
    page_title="PMO BTS 4G Meeting Assistant",
    page_icon="📅",
    layout="wide"
)

st.title("⚡ PMO BTS 4G Meeting Assistant")
st.write("Silakan unggah file Excel Jadwal Anda di panel kiri untuk memulai. Sistem akan mencari template otomatis.")

DAYS_INDO = {
    0: "Senin", 1: "Selasa", 2: "Rabu", 3: "Kamis", 4: "Jumat", 5: "Sabtu", 6: "Minggu"
}
MONTHS_INDO = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 5: "Mei", 6: "Juni",
    7: "Juli", 8: "Agustus", 9: "September", 10: "Oktober", 11: "November", 12: "Desember"
}

def format_date_indo_full(dt_val):
    if isinstance(dt_val, (datetime.datetime, datetime.date)):
        d_name = DAYS_INDO[dt_val.weekday()]
        m_name = MONTHS_INDO[dt_val.month]
        return f"{d_name}, {dt_val.day} {m_name} {dt_val.year}"
    return str(dt_val) if pd.notna(dt_val) else "-"

def format_date_indo_short(dt_val):
    if isinstance(dt_val, (datetime.datetime, datetime.date)):
        m_name = MONTHS_INDO[dt_val.month]
        return f"{dt_val.day} {m_name} {dt_val.year}"
    return str(dt_val) if pd.notna(dt_val) else "-"

def sanitize_filename(name):
    return re.sub(r'[\\/*?:"<>|]', "", str(name))[:50]

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_hex))
    cell._tc.get_or_add_tcPr().append(shading_elm)

# ==============================================================================
# FUNGSI WORDING WA
# ==============================================================================
def build_wording_undangan(tgl, tempat, waktu, agenda, zoom_link, passcode, peserta):
    peserta_list = [p.strip() for p in str(peserta).split(",") if p.strip() and str(p).lower() != 'nan']
    peserta_formatted = "\n".join([f"{i+1}. {p}" for i, p in enumerate(peserta_list)])
    
    zoom_block = f"Join Zoom Meeting\n{zoom_link}\n\n" if zoom_link and str(zoom_link).strip() not in ["-", "", "nan"] else ""
    passcode_block = f"{passcode}\n\n" if passcode and str(passcode).strip() not in ["-", "", "nan"] else ""

    return f"""Kepada Yth Bapak/Ibu, 

Berikut kami sampaikan undangan Meeting, yang akan dilaksanakan pada:

Hari, Tanggal : {tgl}
Tempat         : {tempat}
Waktu          : {waktu}
Agenda         : {agenda}

{zoom_block}{passcode_block}Peserta: 
{peserta_formatted}

Untuk efektivitas waktu diharapkan kepada para undangan untuk dapat hadir tepat waktu.
Terima kasih"""

def build_wording_reminder(tgl, tempat, waktu, agenda, zoom_link, passcode):
    zoom_block = f"🔗 *Join Zoom Meeting:*\n{zoom_link}\n\n" if zoom_link and str(zoom_link).strip() not in ["-", "", "nan"] else ""
    passcode_block = f"{passcode}\n\n" if passcode and str(passcode).strip() not in ["-", "", "nan"] else ""

    return f"""🔔 *[REMINDER]* Undangan Meeting Koordinasi

Yth. Rekan-rekan Tim Peserta Rapat,

Mengingatkan kembali pelaksanaan agenda rapat berikut:
📌 *Agenda:* {agenda}
📅 *Hari/Tgl:* {tgl}
⏰ *Waktu:* {waktu}
🏢 *Tempat:* {tempat}

{zoom_block}{passcode_block}Dimohon kesediaannya untuk bergabung tepat waktu demi kelancaran pembahasan. Terima kasih 🙏"""

# ==============================================================================
# FUNGSI SMART REPLACE UNTUK DOKUMEN WORD (MENGGUNAKAN TEMPLATE LOKAL)
# ==============================================================================
def process_word_template(template_path, agenda, waktu, peserta, tgl_str, tgl_surat, nomor_surat, jenis_kegiatan):
    doc = docx.Document(template_path)
    
    OLD_NOMOR = "2302/BAKTI.31.3/UM.01.01/10/2026"
    OLD_TANGGAL_1 = "23 Oktober 2026"
    OLD_TANGGAL_2 = "23 Oktober2026"
    
    # 1. Ganti Tanggal Surat & Nomor Surat
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if OLD_NOMOR in cell.text:
                    cell.text = cell.text.replace(OLD_NOMOR, nomor_surat)
                if OLD_TANGGAL_1 in cell.text:
                    cell.text = cell.text.replace(OLD_TANGGAL_1, tgl_surat)
                if OLD_TANGGAL_2 in cell.text:
                    cell.text = cell.text.replace(OLD_TANGGAL_2, tgl_surat)

    for p in doc.paragraphs:
        if OLD_NOMOR in p.text:
            p.text = p.text.replace(OLD_NOMOR, nomor_surat)
        if OLD_TANGGAL_1 in p.text:
            p.text = p.text.replace(OLD_TANGGAL_1, tgl_surat)
        if OLD_TANGGAL_2 in p.text:
            p.text = p.text.replace(OLD_TANGGAL_2, tgl_surat)

    # 2. Modifikasi Paragraf Body
    for p in doc.paragraphs:
        if "Sehubungan dengan upaya menjaga performa" in p.text:
            p.clear()
            p.add_run(f"Sehubungan dengan upaya menjaga performa, keandalan, serta ketersediaan jaringan {jenis_kegiatan}, kami mengundang Bapak/Ibu untuk hadir dalam agenda rutin ")
            
            run_agenda = p.add_run(agenda)
            run_agenda.bold = True
            
            p.add_run(" yang akan dilaksanakan pada :")

    # 3. Update Tabel Detail Halaman Depan
    for t in doc.tables:
        if len(t.rows) >= 2 and "Hari/Tanggal" in t.rows[0].cells[0].text:
            t.rows[0].cells[2].text = tgl_str
            t.rows[1].cells[2].text = waktu
            if len(t.rows) > 3:
                t.rows[3].cells[2].text = agenda

    # 4. Modifikasi Lampiran 1
    konsorsium_dihapus = [
        "Konsorsium IBS – ZTE", "Konsorsium FH – MTD", "Konsorsium BKU – Huawei", 
        "Konsorsium Hypernet – Ceragon", "PT. Lindu Putra Utama", "PT Lindu Putra Utama",
        "PT Alita Praya Mitra", "Konsorsium BKU – Langgeng", "Konsorsium SCI – NP"
    ]
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt in konsorsium_dihapus:
            p.clear()

    # 5. Modifikasi Lampiran 2 (Tabel Rundown dengan Layout Khusus)
    tabel_rundown = None
    for t in doc.tables:
        if len(t.rows) > 1 and "Rundown Acara" in t.rows[0].cells[0].text:
            tabel_rundown = t
            break

    if tabel_rundown:
        while len(tabel_rundown.rows) > 2:
            tbl = tabel_rundown._tbl
            tr = tabel_rundown.rows[-1]._tr
            tbl.remove(tr)
        
        # Tambahkan Baris Tanggal (Merged Cell dengan Background Biru)
        row_date = tabel_rundown.add_row()
        a = row_date.cells[0]
        b = row_date.cells[-1]
        a.merge(b) 
        
        a.text = tgl_str
        p_date = a.paragraphs[0]
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_date.runs:
            p_date.runs[0].bold = True
            p_date.runs[0].font.name = "Arial"
        set_cell_background(a, "DDEBF7") 
        
        # Tambahkan Baris Data Rapat
        row_data = tabel_rundown.add_row()
        row_data.cells[0].text = "1"
        row_data.cells[1].text = agenda
        row_data.cells[2].text = waktu
        
        peserta_list = [p.strip() for p in str(peserta).split(',') if p.strip() and str(p).lower() != 'nan']
        peserta_numbered = "\n".join([f"{i+1}.\t{p}" for i, p in enumerate(peserta_list)])
        row_data.cells[3].text = peserta_numbered

    return doc

# ==============================================================================
# ALUR UI APLIKASI
# ==============================================================================

# Cek ketersediaan template di folder lokal
local_templates = glob.glob("*.docx")
template_path = local_templates[0] if local_templates else None

st.sidebar.header("📁 Panel File")
uploaded_excel = st.sidebar.file_uploader("Upload Excel Jadwal (.xlsx)", type=["xlsx"])

if template_path:
    st.sidebar.success(f"✅ Template Terdeteksi: `{template_path}`")
else:
    st.sidebar.error("❌ Template Word (.docx) tidak ditemukan! Pastikan file template berada di folder yang sama dengan aplikasi.")

if uploaded_excel and template_path:
    try:
        xl = pd.ExcelFile(uploaded_excel)
        sheet_choice = st.sidebar.selectbox("Pilih Sheet Bulan:", xl.sheet_names)
        
        df = pd.read_excel(uploaded_excel, sheet_name=sheet_choice).dropna(how="all")
        
        # Identifikasi Kolom Utama (Fleksibel Case-Insensitive)
        cols_lower = [str(c).strip().lower() for c in df.columns]
        
        def find_col(keywords):
            for i, c in enumerate(cols_lower):
                if all(k in c for k in keywords):
                    return df.columns[i]
            return None
        
        col_tgl = find_col(["tanggal"]) or df.columns[0]
        col_waktu = find_col(["waktu"]) or df.columns[2]
        col_lokasi = find_col(["lokasi"]) or df.columns[3]
        col_agenda = find_col(["agenda"]) or df.columns[5]
        col_peserta = find_col(["attendance"]) or (find_col(["peserta"]) or df.columns[6])
        col_zoom = find_col(["zoom"]) or find_col(["link"])
        
        # Ini yang krusial untuk memperbaiki deteksi Nomor Surat dan Jenis Kegiatan
        col_nomor = find_col(["nomor", "surat"]) or find_col(["nomor"])
        col_jenis = find_col(["jenis"]) or find_col(["kegiatan"])

        df = df[df[col_agenda].notna()]
        total_agenda = len(df)
        
        st.subheader(f"📊 Data Jadwal Terbaca: [{sheet_choice}] (Total {total_agenda} Agenda)")
        st.dataframe(df.head(10), use_container_width=True)

        st.divider()
        st.subheader(f"📄 Export Batch ke Format Word (.docx)")
        
        if st.button(f"🚀 Generate {total_agenda} File Word (.zip)"):
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
                for idx, row in df.iterrows():
                    raw_tgl = row.get(col_tgl)
                    val_tgl_full = format_date_indo_full(raw_tgl)
                    val_tgl_short = format_date_indo_short(raw_tgl)
                    
                    val_waktu = str(row.get(col_waktu, "-"))
                    val_agenda = str(row.get(col_agenda, "-"))
                    val_peserta = str(row.get(col_peserta, "-"))
                    
                    # Logika Penarikan Nomor Surat 
                    raw_nomor = row.get(col_nomor) if col_nomor else None
                    val_nomor = str(raw_nomor).strip() if pd.notna(raw_nomor) and str(raw_nomor).strip() not in ["", "nan"] else "2302/BAKTI.31.3/UM.01.01/10/2026"
                    
                    # Logika Penarikan Jenis Kegiatan
                    raw_jenis = row.get(col_jenis) if col_jenis else None
                    val_jenis = str(raw_jenis).strip() if pd.notna(raw_jenis) and str(raw_jenis).strip() not in ["", "nan"] else "BTS 4G"
                    
                    doc = process_word_template(
                        template_path=template_path,
                        agenda=val_agenda,
                        waktu=val_waktu,
                        peserta=val_peserta,
                        tgl_str=val_tgl_full,
                        tgl_surat=val_tgl_short,
                        nomor_surat=val_nomor,
                        jenis_kegiatan=val_jenis
                    )

                    doc_bio = io.BytesIO()
                    doc.save(doc_bio)
                    
                    safe_filename = sanitize_filename(val_agenda)
                    file_name = f"Undangan_{idx+1}_{safe_filename}.docx"
                    zip_file.writestr(file_name, doc_bio.getvalue())

            st.success(f"✅ Berhasil membuat {total_agenda} file Word ke dalam 1 file ZIP!")
            st.download_button(
                label="📥 Unduh Semua File Word (.zip)",
                data=zip_buffer.getvalue(),
                file_name=f"Kumpulan_Undangan_Update_{sheet_choice}.zip",
                mime="application/zip"
            )

        st.divider()
        st.subheader("📲 Generator Wording WhatsApp")
        
        agenda_options = [f"Baris {idx+1}: {row.get(col_agenda, 'Tanpa Agenda')}" for idx, row in df.iterrows()]
        selected_idx = st.selectbox("Pilih Agenda Rapat:", range(len(agenda_options)), format_func=lambda x: agenda_options[x])

        selected_row = df.iloc[selected_idx]
        val_tgl = format_date_indo_full(selected_row.get(col_tgl))
        val_waktu = str(selected_row.get(col_waktu, "-"))
        val_lokasi = str(selected_row.get(col_lokasi, "-"))
        val_agenda = str(selected_row.get(col_agenda, "-"))
        val_peserta = str(selected_row.get(col_peserta, "-"))
        
        # Ambil zoom link 
        raw_zoom = selected_row.get(col_zoom) if col_zoom else ""
        val_zoom = str(raw_zoom) if pd.notna(raw_zoom) and str(raw_zoom) != "nan" else ""

        st.markdown("##### 🔧 Kustomisasi Akses Virtual (Khusus WhatsApp)")
        col_in1, col_in2 = st.columns(2)
        with col_in1:
            custom_zoom = st.text_input("Tautan / Link Zoom Meeting:", value=val_zoom)
        with col_in2:
            custom_passcode = st.text_input("Meeting ID & Passcode (Opsional):", value="")

        wording_resmi = build_wording_undangan(val_tgl, val_lokasi, val_waktu, val_agenda, custom_zoom, custom_passcode, val_peserta)
        wording_reminder = build_wording_reminder(val_tgl, val_lokasi, val_waktu, val_agenda, custom_zoom, custom_passcode)

        tab1, tab2 = st.tabs(["📋 Format 1: Undangan Lengkap", "🔔 Format 2: Reminder Cepat"])

        with tab1:
            st.text_area("Salin Wording Undangan Resmi (Blok & Ctrl + C):", value=wording_resmi, height=300)

        with tab2:
            st.text_area("Salin Wording Reminder (Blok & Ctrl + C):", value=wording_reminder, height=280)

    except Exception as e:
        st.error(f"Terjadi kesalahan saat memproses file: {e}")